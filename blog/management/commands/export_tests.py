import os
from django.core.management.base import BaseCommand
from django.conf import settings
from docx import Document
from docx.shared import Pt, RGBColor


class Command(BaseCommand):
    help = 'Экспортирует код тестов из файла blog/tests.py в документ Word'

    def handle(self, *args, **options):
        tests_path = os.path.join(settings.BASE_DIR, 'blog', 'tests.py')

        if not os.path.exists(tests_path):
            self.stderr.write(f"Файл {tests_path} не найден!")
            return

        with open(tests_path, 'r', encoding='utf-8') as f:
            code_content = f.read()

        doc = Document()
        doc.add_heading('Django Unit & Integration Tests', level=0)
        doc.add_paragraph('Автоматический экспорт тестов из файла blog/tests.py:')

        doc.add_heading('Код тестов (blog/tests.py)', level=1)

        p = doc.add_paragraph()
        run = p.add_run(code_content)
        run.font.name = 'Courier New'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(30, 30, 30)

        output_path = os.path.join(settings.BASE_DIR, 'django_tests_report.docx')
        doc.save(output_path)

        self.stdout.write(self.style.SUCCESS(f'Файл успешно создан: {output_path}'))